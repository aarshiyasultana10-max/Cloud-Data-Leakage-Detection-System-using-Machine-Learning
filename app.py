from flask import Flask, render_template, request, redirect, session, jsonify, send_file
import sqlite3, os, re, io, csv, hashlib
from datetime import datetime
from cryptography.fernet import Fernet
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.svm import SVC
from sklearn.model_selection import train_test_split
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, confusion_matrix
from functools import wraps

try:
    from fpdf import FPDF
    PDF_AVAILABLE = True
except:
    PDF_AVAILABLE = False

try:
    import spacy
    nlp = spacy.load("en_core_web_sm")
    SPACY_AVAILABLE = True
except:
    SPACY_AVAILABLE = False

app = Flask(__name__)
app.secret_key = "dlpshield_secret_2024"
UPLOAD_FOLDER = "uploads"
QUARANTINE_FOLDER = "quarantine"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(QUARANTINE_FOLDER, exist_ok=True)

# ── Expanded ML Training Data ──────────────────────────
train_texts = [
    "email test@gmail.com","hello world","password 1234","normal text",
    "credit card 1234567890123456","my aadhaar is 1234 5678 9012",
    "pan card ABCDE1234F","normal project report","meeting notes today",
    "user password secret123","account number 9876543210","safe document",
    "confidential data leaked","resume objective statement","invoice number 12345",
    "contact me at john@company.com","my phone is 9876543210","send to finance team",
    "ssn 123-45-6789","visa card 4532 1234 5678 9012","passport A1234567",
    "bank account 123456789012","salary details confidential","private key exposed",
    "api_key sk_live_abcdef12345","token eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9",
    "quarterly report summary","project meeting agenda","team lunch thursday",
    "server password root123","database credentials mysql","env variable SECRET_KEY",
    "internal memo for hr","public announcement newsletter","conference schedule 2025",
    "customer ssn 987-65-4321","employee pan PQRST5678U","voter id ABC1234567",
    "driving license DL123456789012","health record diagnosis","insurance policy number",
    "regular office document","training material slides","product roadmap planning",
    "user data export email phone","private key rsa 2048","certificate authority file",
    "normal business letter","marketing campaign draft","code review comments",
    "aadhaar linked mobile 9123456780","credit score report leaked","bank statement pdf",
    "personal details name dob address","public blog post article","open source readme",
    "DB_PASSWORD admin123","aws access key AKIAIOSFODNN7EXAMPLE","git token ghp_abc",
    "class notes lecture slides","homework assignment submission","exam schedule college",
]
train_labels = [
    "leak","safe","leak","safe","leak","leak","leak","safe","safe","leak","leak","safe","leak","safe","safe",
    "safe","leak","safe","leak","leak","leak","leak","leak","leak","leak","leak",
    "safe","safe","safe","leak","leak","leak","safe","safe","safe",
    "leak","leak","leak","leak","leak","leak","safe","safe","safe",
    "leak","leak","leak","leak","safe","safe","leak","leak","leak",
    "leak","safe","safe","leak","leak","leak","safe","safe","safe",
]

vectorizer = TfidfVectorizer(ngram_range=(1, 2), max_features=500)
X = vectorizer.fit_transform(train_texts)
y = train_labels

X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

model_rf  = RandomForestClassifier(n_estimators=150, random_state=42)
model_lr  = LogisticRegression(max_iter=1000, random_state=42)
model_svm = SVC(probability=True, kernel='rbf', random_state=42)

for m in [model_rf, model_lr, model_svm]:
    m.fit(X_train, y_train)

model = model_rf  # primary

def get_model_metrics():
    metrics = {}
    for name, m in [("Random Forest", model_rf), ("Logistic Regression", model_lr), ("SVM", model_svm)]:
        preds = m.predict(X_test)
        try:
            cm = confusion_matrix(y_test, preds, labels=["leak","safe"]).tolist()
        except:
            cm = [[0,0],[0,0]]
        metrics[name] = {
            "accuracy":  round(accuracy_score(y_test, preds) * 100, 1),
            "precision": round(precision_score(y_test, preds, pos_label="leak", zero_division=0) * 100, 1),
            "recall":    round(recall_score(y_test, preds, pos_label="leak", zero_division=0) * 100, 1),
            "f1":        round(f1_score(y_test, preds, pos_label="leak", zero_division=0) * 100, 1),
            "cm": cm
        }
    return metrics

# ── DB ────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    return conn

def hash_password(pw):
    return hashlib.sha256(pw.encode()).hexdigest()

def init_db():
    conn = get_db(); c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS users(
        id INTEGER PRIMARY KEY, username TEXT UNIQUE, password TEXT, role TEXT,
        failed_attempts INTEGER DEFAULT 0, locked INTEGER DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS files(
        id INTEGER PRIMARY KEY, filename TEXT, owner TEXT, status TEXT,
        rule_detections TEXT, ml_prediction TEXT, confidence REAL,
        file_size INTEGER DEFAULT 0, risk_score INTEGER DEFAULT 0,
        ner_entities TEXT DEFAULT '', quarantined INTEGER DEFAULT 0,
        uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS attack_logs(
        id INTEGER PRIMARY KEY, event TEXT, ip TEXT, username TEXT,
        timestamp TEXT DEFAULT CURRENT_TIMESTAMP, severity TEXT DEFAULT 'medium')""")
    c.execute("""CREATE TABLE IF NOT EXISTS honeypot_log(
        id INTEGER PRIMARY KEY, ip TEXT, user TEXT, timestamp TEXT DEFAULT CURRENT_TIMESTAMP)""")
    for col, defn in [("risk_score","INTEGER DEFAULT 0"),("ner_entities","TEXT DEFAULT ''"),("quarantined","INTEGER DEFAULT 0")]:
        try: c.execute(f"ALTER TABLE files ADD COLUMN {col} {defn}")
        except: pass
    try:
        c.execute("INSERT INTO users(username,password,role) VALUES (?,?,?)",
                  ("admin", hash_password("admin123"), "Admin"))
    except: pass
    conn.commit(); conn.close()

init_db()

# ── Detection ─────────────────────────────────────────
def rule_detect(text):
    patterns = {
        "Email":        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        "Phone":        r"\b[6-9]\d{9}\b",
        "Credit Card":  r"\b\d{4}[\s-]?\d{4}[\s-]?\d{4}[\s-]?\d{4}\b",
        "Password":     r"\b(password|passwd|pwd)\s*[=:\s]+\S+",
        "Aadhaar":      r"\b\d{4}\s?\d{4}\s?\d{4}\b",
        "PAN Card":     r"\b[A-Z]{5}[0-9]{4}[A-Z]\b",
        "Passport":     r"\b[A-PR-WY][1-9]\d{7}\b",
        "SSN":          r"\b\d{3}-\d{2}-\d{4}\b",
        "Bank Account": r"\b\d{9,18}\b",
        "API Key":      r"\b(sk_live_|sk_test_|AKIA|ghp_|glpat-)[A-Za-z0-9_\-]{10,}",
        "JWT Token":    r"\beyJ[A-Za-z0-9_\-]+\.[A-Za-z0-9_\-]+\.[A-Za-z0-9_\-]+",
        "Voter ID":     r"\b[A-Z]{3}[0-9]{7}\b",
        "Env Secret":   r"\b(SECRET_KEY|DB_PASSWORD|API_SECRET|ACCESS_TOKEN)\s*=\s*\S+",
        "Private Key":  r"-----BEGIN (RSA |EC )?PRIVATE KEY-----",
    }
    return [k for k, v in patterns.items() if re.search(v, text, re.I)]

def ml_detect(text):
    Xv = vectorizer.transform([text])
    pred = model.predict(Xv)[0]
    proba = model.predict_proba(Xv)[0]
    classes = list(model.classes_)
    confidence = float(proba[classes.index(pred)]) * 100
    return pred, round(confidence, 1)

def ner_detect(text):
    if not SPACY_AVAILABLE: return []
    doc = nlp(text[:5000])
    found = []
    for ent in doc.ents:
        if ent.label_ in ("PERSON","ORG","GPE","MONEY","CARDINAL"):
            found.append(f"{ent.text} [{ent.label_}]")
    return list(set(found))[:10]

def compute_risk_score(rule_hits, ml_pred, confidence, file_size_bytes):
    score = min(len(rule_hits) * 15, 45)
    if ml_pred == "leak":
        score += int((confidence / 100) * 35)
    size_kb = file_size_bytes / 1024
    if size_kb > 500: score += 10
    elif size_kb > 100: score += 5
    return min(int(score), 100)

def extract_text_from_file(data, filename):
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    try:
        if ext == 'pdf':
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(data))
                return " ".join(p.extract_text() or "" for p in reader.pages)
            except: pass
        elif ext == 'docx':
            try:
                import docx as docxlib
                doc = docxlib.Document(io.BytesIO(data))
                return " ".join(p.text for p in doc.paragraphs)
            except: pass
        return data.decode("utf-8", errors="ignore")
    except: return ""

def log_attack(event, severity="medium", username="unknown"):
    try:
        conn = get_db()
        conn.execute("INSERT INTO attack_logs(event,ip,username,severity) VALUES(?,?,?,?)",
                     (event, request.remote_addr, username, severity))
        conn.commit(); conn.close()
    except: pass

# ── Decorators ────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user" not in session: return redirect("/")
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if session.get("role") != "Admin": return redirect("/dashboard")
        return f(*args, **kwargs)
    return decorated

# ── Routes ────────────────────────────────────────────
@app.route("/", methods=["GET","POST"])
def login():
    error = None
    if request.method == "POST":
        u = request.form["username"]; p = request.form["password"]
        conn = get_db()
        user = conn.execute("SELECT * FROM users WHERE username=?", (u,)).fetchone()
        if user and user["locked"]:
            error = "Account locked. Contact admin."
            log_attack(f"Locked account attempt: {u}", "high", u); conn.close()
        elif user and user["password"] == hash_password(p):
            conn.execute("UPDATE users SET failed_attempts=0 WHERE username=?", (u,))
            conn.commit(); conn.close()
            session["user"] = u; session["role"] = user["role"]
            return redirect("/dashboard")
        else:
            attempts = (user["failed_attempts"] if user else 0) + 1
            if user:
                locked = 1 if attempts >= 3 else 0
                conn.execute("UPDATE users SET failed_attempts=?,locked=? WHERE username=?", (attempts,locked,u))
                conn.commit()
                if locked:
                    log_attack(f"Account locked (brute force): {u}", "critical", u)
                    error = "Account locked after 3 failed attempts!"
                else:
                    error = f"Invalid credentials. {3-attempts} attempt(s) left."
            else: error = "User not found."
            log_attack(f"Failed login: {u}", "medium", u); conn.close()
    return render_template("login.html", error=error)

@app.route("/register", methods=["GET","POST"])
def register():
    error = None
    if request.method == "POST":
        u = request.form["username"]; p = request.form["password"]; role = request.form["role"]
        if len(p) < 8: error = "Password must be at least 8 characters."
        elif not re.search(r'[A-Z]', p): error = "Password must contain at least one uppercase letter."
        elif not re.search(r'[0-9]', p): error = "Password must contain at least one number."
        else:
            try:
                conn = get_db()
                conn.execute("INSERT INTO users(username,password,role) VALUES(?,?,?)",(u, hash_password(p), role))
                conn.commit(); conn.close()
                return redirect("/")
            except: error = "Username already exists."
    return render_template("register.html", error=error)

@app.route("/logout")
def logout():
    session.clear(); return redirect("/")

@app.route("/dashboard")
@login_required
def dashboard():
    conn = get_db()
    total   = conn.execute("SELECT COUNT(*) FROM files").fetchone()[0]
    leaks   = conn.execute("SELECT COUNT(*) FROM files WHERE status='Leak'").fetchone()[0]
    safe    = conn.execute("SELECT COUNT(*) FROM files WHERE status='Safe'").fetchone()[0]
    users   = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
    attacks = conn.execute("SELECT COUNT(*) FROM attack_logs").fetchone()[0]
    quarantined = conn.execute("SELECT COUNT(*) FROM files WHERE quarantined=1").fetchone()[0]
    recent  = conn.execute("SELECT * FROM files ORDER BY id DESC LIMIT 5").fetchall()
    activity = conn.execute(
        "SELECT DATE(uploaded_at) as day, COUNT(*) as cnt FROM files WHERE uploaded_at >= date('now','-30 days') GROUP BY day"
    ).fetchall()
    conn.close()
    activity_map = {row['day']: row['cnt'] for row in activity}
    return render_template("dashboard.html", role=session.get("role"), user=session.get("user"),
        total=total, leaks=leaks, safe=safe, users=users, attacks=attacks,
        quarantined=quarantined, recent=recent, activity_map=activity_map)

@app.route("/upload", methods=["GET","POST"])
@login_required
def upload():
    result=None; ml_pred=None; confidence=None; status=None; fname=None
    risk_score=0; ner_entities=[]; quarantined=False
    if request.method == "POST":
        file = request.files["file"]; data = file.read(); fname = file.filename; size = len(data)
        key = Fernet.generate_key(); encrypted_data = Fernet(key).encrypt(data)
        text = extract_text_from_file(data, fname)
        result = rule_detect(text)
        ml_pred, confidence = ml_detect(text)
        ner_entities = ner_detect(text)
        status = "Leak" if result or ml_pred == "leak" else "Safe"
        risk_score = compute_risk_score(result, ml_pred, confidence, size)
        if status == "Leak":
            open(os.path.join(QUARANTINE_FOLDER, fname), "wb").write(encrypted_data)
            quarantined = True
            log_attack(f"Leak detected & quarantined: {fname} (risk={risk_score})", "high", session["user"])
        else:
            open(os.path.join(UPLOAD_FOLDER, fname), "wb").write(encrypted_data)
        ner_str = "; ".join(ner_entities) if ner_entities else ""
        conn = get_db()
        conn.execute(
            "INSERT INTO files(filename,owner,status,rule_detections,ml_prediction,confidence,file_size,risk_score,ner_entities,quarantined) VALUES(?,?,?,?,?,?,?,?,?,?)",
            (fname, session["user"], status, ",".join(result), ml_pred, confidence, size, risk_score, ner_str, 1 if quarantined else 0))
        conn.commit(); conn.close()
    return render_template("upload.html", result=result, ml=ml_pred, confidence=confidence,
                           status=status, fname=fname, risk_score=risk_score,
                           ner_entities=ner_entities, quarantined=quarantined,
                           spacy_available=SPACY_AVAILABLE)

@app.route("/scan-text", methods=["GET","POST"])
@login_required
def scan_text():
    result=None; ml_pred=None; confidence=None; status=None; risk_score=0; ner_entities=[]; text_input=""
    if request.method == "POST":
        text_input = request.form.get("text_input","")
        if text_input.strip():
            result = rule_detect(text_input)
            ml_pred, confidence = ml_detect(text_input)
            ner_entities = ner_detect(text_input)
            status = "Leak" if result or ml_pred == "leak" else "Safe"
            risk_score = compute_risk_score(result, ml_pred, confidence, len(text_input.encode()))
            if status == "Leak":
                log_attack(f"Text scan leak detected (risk={risk_score})", "high", session["user"])
    return render_template("scan_text.html", result=result, ml=ml_pred, confidence=confidence,
                           status=status, risk_score=risk_score, ner_entities=ner_entities,
                           text_input=text_input, spacy_available=SPACY_AVAILABLE)

@app.route("/api/scan", methods=["POST"])
@login_required
def api_scan():
    data = request.get_json(silent=True)
    if not data or "text" not in data:
        return jsonify(error="Provide JSON body with 'text' field"), 400
    text = data["text"]
    rules = rule_detect(text); ml_pred, confidence = ml_detect(text)
    ner_entities = ner_detect(text)
    status = "leak" if rules or ml_pred == "leak" else "safe"
    risk = compute_risk_score(rules, ml_pred, confidence, len(text.encode()))
    return jsonify(status=status, rules=rules, ml_prediction=ml_pred,
                   confidence=confidence, risk_score=risk, ner_entities=ner_entities)

@app.route("/api/stats")
@login_required
def api_stats():
    conn = get_db()
    total   = conn.execute("SELECT COUNT(*) FROM files").fetchone()[0]
    leaks   = conn.execute("SELECT COUNT(*) FROM files WHERE status='Leak'").fetchone()[0]
    safe    = conn.execute("SELECT COUNT(*) FROM files WHERE status='Safe'").fetchone()[0]
    attacks = conn.execute("SELECT COUNT(*) FROM attack_logs").fetchone()[0]
    quarantined = conn.execute("SELECT COUNT(*) FROM files WHERE quarantined=1").fetchone()[0]
    conn.close()
    return jsonify(total=total, leaks=leaks, safe=safe, attacks=attacks, quarantined=quarantined)

@app.route("/api/docs")
@login_required
def api_docs():
    return render_template("api_docs.html")

@app.route("/history")
@login_required
def history():
    conn = get_db()
    if session.get("role") == "Admin":
        files = conn.execute("SELECT * FROM files ORDER BY id DESC").fetchall()
    else:
        files = conn.execute("SELECT * FROM files WHERE owner=? ORDER BY id DESC", (session["user"],)).fetchall()
    conn.close()
    return render_template("history.html", files=files)

@app.route("/graph")
@login_required
def graph():
    conn = get_db()
    status_data   = conn.execute("SELECT status, COUNT(*) FROM files GROUP BY status").fetchall()
    daily_data    = conn.execute("SELECT DATE(uploaded_at), COUNT(*) FROM files GROUP BY DATE(uploaded_at) ORDER BY DATE(uploaded_at)").fetchall()
    filetype_data = conn.execute("SELECT LOWER(SUBSTR(filename, INSTR(filename,'.')+1)) as ext, COUNT(*) FROM files GROUP BY ext").fetchall()
    risk_data     = conn.execute("SELECT risk_score FROM files WHERE risk_score IS NOT NULL ORDER BY uploaded_at DESC LIMIT 30").fetchall()
    conn.close()
    return render_template("graph.html",
        labels=[r[0] for r in status_data], values=[r[1] for r in status_data],
        daily_labels=[r[0] for r in daily_data], daily_values=[r[1] for r in daily_data],
        ft_labels=[r[0] or 'other' for r in filetype_data], ft_values=[r[1] for r in filetype_data],
        risk_values=[r[0] for r in risk_data])

@app.route("/admin")
@login_required
@admin_required
def admin():
    conn = get_db()
    users = conn.execute("SELECT * FROM users ORDER BY id DESC").fetchall()
    files = conn.execute("SELECT * FROM files ORDER BY id DESC").fetchall()
    logs  = conn.execute("SELECT * FROM attack_logs ORDER BY id DESC LIMIT 50").fetchall()
    quarantined = conn.execute("SELECT * FROM files WHERE quarantined=1 ORDER BY id DESC").fetchall()
    model_metrics = get_model_metrics()
    conn.close()
    return render_template("admin.html", users=users, files=files, logs=logs,
                           quarantined=quarantined, model_metrics=model_metrics)

@app.route("/admin/delete_user/<int:uid>")
@login_required
@admin_required
def delete_user(uid):
    conn = get_db()
    conn.execute("DELETE FROM users WHERE id=?", (uid,))
    conn.commit(); conn.close()
    return redirect("/admin")

@app.route("/admin/unlock_user/<int:uid>")
@login_required
@admin_required
def unlock_user(uid):
    conn = get_db()
    conn.execute("UPDATE users SET locked=0,failed_attempts=0 WHERE id=?", (uid,))
    conn.commit(); conn.close()
    return redirect("/admin")

@app.route("/admin/release_quarantine/<int:fid>")
@login_required
@admin_required
def release_quarantine(fid):
    conn = get_db()
    f = conn.execute("SELECT * FROM files WHERE id=?", (fid,)).fetchone()
    if f:
        q_path = os.path.join(QUARANTINE_FOLDER, f["filename"])
        safe_path = os.path.join(UPLOAD_FOLDER, f["filename"])
        if os.path.exists(q_path):
            os.rename(q_path, safe_path)
        conn.execute("UPDATE files SET quarantined=0 WHERE id=?", (fid,))
        conn.commit()
        log_attack(f"Quarantine released by admin: {f['filename']}", "medium", session["user"])
    conn.close()
    return redirect("/admin")

@app.route("/attack_log")
@login_required
def attack_log():
    severity_filter = request.args.get("severity","")
    search = request.args.get("search","")
    conn = get_db()
    query = "SELECT * FROM attack_logs WHERE 1=1"
    params = []
    if severity_filter: query += " AND severity=?"; params.append(severity_filter)
    if search:
        query += " AND (event LIKE ? OR username LIKE ?)"
        params.extend([f"%{search}%", f"%{search}%"])
    query += " ORDER BY id DESC"
    logs     = conn.execute(query, params).fetchall()
    total    = conn.execute("SELECT COUNT(*) FROM attack_logs").fetchone()[0]
    critical = conn.execute("SELECT COUNT(*) FROM attack_logs WHERE severity='critical'").fetchone()[0]
    high     = conn.execute("SELECT COUNT(*) FROM attack_logs WHERE severity='high'").fetchone()[0]
    conn.close()
    return render_template("attack_log.html", logs=logs, total=total, critical=critical, high=high,
                           severity_filter=severity_filter, search=search)

@app.route("/attacker")
@login_required
def attacker():
    log_attack("Unauthorized area accessed: /attacker", "high", session.get("user","unknown"))
    return render_template("attacker.html")

@app.route("/honeypot")
def honeypot():
    conn = get_db()
    conn.execute("INSERT INTO honeypot_log(ip,user) VALUES(?,?)",
                 (request.remote_addr, session.get("user","anonymous")))
    conn.commit()
    log_attack("HONEYPOT TRIGGERED — Fake admin link clicked", "critical", session.get("user","anonymous"))
    conn.close()
    return render_template("honeypot.html")

@app.route("/export/pdf")
@login_required
def export_pdf():
    if not PDF_AVAILABLE:
        return "fpdf2 not installed. Run: pip install fpdf2", 400
    conn = get_db()
    if session.get("role") == "Admin":
        files = conn.execute("SELECT * FROM files ORDER BY id DESC").fetchall()
    else:
        files = conn.execute("SELECT * FROM files WHERE owner=? ORDER BY id DESC", (session["user"],)).fetchall()
    stats = conn.execute("SELECT status, COUNT(*) FROM files GROUP BY status").fetchall()
    conn.close()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica","B",20); pdf.set_text_color(0,0,0)
    pdf.cell(0,14,"DLP Shield - Security Report",ln=True,align="C")
    pdf.set_font("Helvetica","",10); pdf.set_text_color(100,100,100)
    pdf.cell(0,6,f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  User: {session['user']}",ln=True,align="C")
    pdf.ln(6); pdf.set_text_color(0,0,0)
    pdf.set_font("Helvetica","B",13); pdf.cell(0,10,"Summary",ln=True)
    pdf.set_font("Helvetica","",10)
    for row in stats: pdf.cell(0,7,f"  {row[0]}: {row[1]} file(s)",ln=True)
    pdf.ln(4); pdf.set_font("Helvetica","B",13); pdf.cell(0,10,"File Scan History",ln=True)
    pdf.set_font("Helvetica","B",9); pdf.set_fill_color(220,220,220)
    for h,w in [("ID",12),("Filename",45),("Owner",25),("Status",20),("Risk",15),("ML",20),("Conf%",16),("Date",30)]:
        pdf.cell(w,7,h,border=1,fill=True)
    pdf.ln(); pdf.set_font("Helvetica","",8)
    for f in files:
        risk = f["risk_score"] if f["risk_score"] else 0
        pdf.cell(12,6,str(f["id"]),border=1); pdf.cell(45,6,str(f["filename"])[:26],border=1)
        pdf.cell(25,6,str(f["owner"]),border=1); pdf.cell(20,6,str(f["status"]),border=1)
        pdf.cell(15,6,str(risk),border=1); pdf.cell(20,6,str(f["ml_prediction"]),border=1)
        pdf.cell(16,6,f"{f['confidence']}%",border=1); pdf.cell(30,6,str(f["uploaded_at"])[:16],border=1)
        pdf.ln()
    buf = io.BytesIO(pdf.output()); buf.seek(0)
    return send_file(buf, as_attachment=True,
        download_name=f"dlp_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
        mimetype="application/pdf")

@app.route("/export/csv")
@login_required
def export_csv():
    conn = get_db()
    if session.get("role") == "Admin":
        files = conn.execute("SELECT * FROM files ORDER BY id DESC").fetchall()
    else:
        files = conn.execute("SELECT * FROM files WHERE owner=? ORDER BY id DESC", (session["user"],)).fetchall()
    conn.close()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ID","Filename","Owner","Status","Rule Detections","ML Prediction","Confidence %","Risk Score","NER Entities","File Size","Quarantined","Uploaded At"])
    for f in files:
        writer.writerow([f["id"],f["filename"],f["owner"],f["status"],f["rule_detections"],
                         f["ml_prediction"],f["confidence"],f["risk_score"],f["ner_entities"],
                         f["file_size"],"Yes" if f["quarantined"] else "No",f["uploaded_at"]])
    output.seek(0)
    return send_file(io.BytesIO(output.getvalue().encode()), as_attachment=True,
        download_name=f"dlp_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
        mimetype="text/csv")

if __name__ == "__main__":
    app.run(debug=True)
